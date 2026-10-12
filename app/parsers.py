from pathlib import Path
import fitz
from docx import Document
from pptx import Presentation


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        doc = fitz.open(path)
        try:
            return "\n\n".join(page.get_text("text") for page in doc)
        finally:
            doc.close()
    if ext == ".docx":
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if ext == ".pptx":
        prs = Presentation(path)
        out = []
        for i, slide in enumerate(prs.slides, start=1):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)
                if shape.has_text_frame is False and getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            slide_text.append(" | ".join(cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                slide_text.append("Notes: " + slide.notes_slide.notes_text_frame.text)
            if slide_text:
                out.append(f"[Slide {i}]\n" + "\n".join(slide_text))
        return "\n\n".join(out)
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError("Unsupported file type.")
